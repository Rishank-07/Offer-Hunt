"""Spreadsheet persistence MCP tool for OFFERHUNT (Milestone M1).

Implements atomic, verified upsert operations into Excel (.xlsx) workbooks with
strict directory confinement, path traversal rejection, and read-back verification.
"""

from __future__ import annotations

import hashlib
import os
import shutil
import tempfile
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, List, Optional, Union

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from agent.contracts import ErrorModel, InternshipRecord


# Default output directory allowed for spreadsheets
DEFAULT_OUTPUT_DIR = Path("artifacts").resolve()

# Columns stored in the spreadsheet
SPREADSHEET_COLUMNS = [
    "record_key",
    "title",
    "company",
    "location",
    "eligibility",
    "stipend",
    "deadline",
    "source_url",
    "source_title",
    "snippet",
    "retrieved_at_utc",
    "source_type",
    "confidence",
    "status",
]


class SpreadsheetSecurityError(ValueError):
    """Raised when an unsafe path, directory traversal, or disallowed extension is detected."""
    pass


class SpreadsheetCorruptError(RuntimeError):
    """Raised when an existing spreadsheet file is corrupted or unreadable."""
    pass


@dataclass
class UpsertResult:
    """Result of an atomic spreadsheet upsert operation."""
    inserted: int
    updated: int
    unchanged: int
    keys: List[str]
    artifact_path: str
    sha256: str


def compute_file_sha256(file_path: Union[str, Path]) -> str:
    """Compute the hex-encoded SHA-256 hash of a file."""
    h = hashlib.sha256()
    with open(file_path, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


def validate_and_resolve_path(
    target_path: Union[str, Path],
    allowed_dir: Union[str, Path] = DEFAULT_OUTPUT_DIR,
) -> Path:
    """Validate that target_path is within allowed_dir and has .xlsx extension.

    Rejects directory traversal and disallowed extensions.
    """
    allowed_dir_resolved = Path(allowed_dir).resolve()

    target = Path(target_path)
    if not target.is_absolute():
        resolved_target = (allowed_dir_resolved / target).resolve()
    else:
        resolved_target = target.resolve()

    # 1. Enforce allowed extension
    if resolved_target.suffix.lower() != ".xlsx":
        raise SpreadsheetSecurityError(
            f"Unsupported spreadsheet extension '{resolved_target.suffix}'. Only '.xlsx' is permitted."
        )

    # 2. Enforce directory confinement (prevent directory traversal)
    try:
        resolved_target.relative_to(allowed_dir_resolved)
    except ValueError:
        raise SpreadsheetSecurityError(
            f"Access denied: path '{target_path}' traverses outside allowed directory '{allowed_dir_resolved}'."
        )

    return resolved_target


def read_spreadsheet_records(file_path: Path) -> tuple[List[str], dict[str, tuple[int, dict[str, Any]]]]:
    """Read existing records from an xlsx workbook.

    Returns:
        (headers, mapping of record_key -> (row_index_1_based, row_dict))
    Raises:
        SpreadsheetCorruptError if workbook cannot be loaded or parsed.
    """
    if not file_path.exists():
        return list(SPREADSHEET_COLUMNS), {}

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as exc:
        raise SpreadsheetCorruptError(
            f"Failed to read existing spreadsheet at {file_path}: {exc}"
        ) from exc

    try:
        ws = wb.active
        if ws is None:
            raise SpreadsheetCorruptError("Workbook does not contain an active sheet.")

        # Read header row
        header_row = [str(cell.value or "").strip() for cell in ws[1]]
        if not header_row or not any(header_row):
            headers = list(SPREADSHEET_COLUMNS)
        else:
            headers = header_row

        if "record_key" not in headers:
            # If empty sheet or headers not established, adopt standard
            headers = list(SPREADSHEET_COLUMNS)

        key_col_idx = headers.index("record_key")

        records: dict[str, tuple[int, dict[str, Any]]] = {}
        for row_idx, row in enumerate(
            ws.iter_rows(min_row=2, max_col=len(headers), values_only=True), start=2
        ):
            if not row or all(v is None for v in row):
                continue
            key_val = row[key_col_idx]
            if key_val is None or str(key_val).strip() == "":
                continue
            record_key = str(key_val).strip()
            row_dict = {
                headers[i]: row[i] if i < len(row) else None
                for i in range(len(headers))
            }
            records[record_key] = (row_idx, row_dict)

        return headers, records
    finally:
        wb.close()


def upsert_records_to_spreadsheet(
    records: Iterable[Union[InternshipRecord, dict[str, Any]]],
    target_path: Union[str, Path] = "offers.xlsx",
    allowed_dir: Union[str, Path] = DEFAULT_OUTPUT_DIR,
) -> UpsertResult:
    """Atomically upsert records into an Excel workbook.

    - Resolves and secures the path against traversal.
    - Never append-only: existing record_keys are updated in-place or marked unchanged.
    - Uses atomic temp-file write, close, and replace.
    - Performs read-back verification before returning.
    """
    resolved_path = validate_and_resolve_path(target_path, allowed_dir=allowed_dir)
    resolved_path.parent.mkdir(parents=True, exist_ok=True)

    # 1. Read existing records if file exists
    headers, existing_records = read_spreadsheet_records(resolved_path)

    # Ensure all standard columns are in headers
    for col in SPREADSHEET_COLUMNS:
        if col not in headers:
            headers.append(col)

    # Convert incoming records to standardized dicts
    input_records: List[dict[str, Any]] = []
    for r in records:
        if isinstance(r, InternshipRecord):
            input_records.append(r.model_dump())
        elif isinstance(r, dict):
            input_records.append(dict(r))
        else:
            raise TypeError(f"Unsupported record type: {type(r)}")

    # Load existing workbook or create new one
    if resolved_path.exists():
        try:
            wb = openpyxl.load_workbook(resolved_path)
            ws = wb.active
        except Exception as exc:
            raise SpreadsheetCorruptError(
                f"Failed to open workbook for editing: {exc}"
            ) from exc
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Offers"
        # Write headers
        ws.append(headers)

    inserted = 0
    updated = 0
    unchanged = 0
    touched_keys: List[str] = []

    for item in input_records:
        r_key = str(item.get("record_key") or "").strip()
        if not r_key:
            raise ValueError("All records must have a non-empty 'record_key'.")

        touched_keys.append(r_key)

        # Standardize string representations for comparison
        def sanitize_val(v: Any) -> Any:
            if v is None:
                return None
            if isinstance(v, (int, float, bool)):
                return v
            return str(v)

        item_row_values = [sanitize_val(item.get(col)) for col in headers]

        if r_key in existing_records:
            row_idx, existing_row_dict = existing_records[r_key]
            # Check if identical
            is_same = True
            for col in headers:
                curr_val = sanitize_val(existing_row_dict.get(col))
                new_val = sanitize_val(item.get(col))
                # For floats (confidence), allow small delta
                if isinstance(curr_val, (int, float)) and isinstance(new_val, (int, float)):
                    if abs(curr_val - new_val) > 1e-6:
                        is_same = False
                        break
                elif str(curr_val or "") != str(new_val or ""):
                    is_same = False
                    break

            if is_same:
                unchanged += 1
            else:
                # Update row in place
                for col_idx, col_name in enumerate(headers, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=item_row_values[col_idx - 1])
                updated += 1
                # Update our in-memory record of the row
                existing_records[r_key] = (row_idx, {col: item_row_values[i] for i, col in enumerate(headers)})
        else:
            # Append new row
            ws.append(item_row_values)
            new_row_idx = ws.max_row
            existing_records[r_key] = (new_row_idx, {col: item_row_values[i] for i, col in enumerate(headers)})
            inserted += 1

    # 2. Atomic Write: write to unique temp file in same directory, close, then replace
    temp_filename = f".tmp_{resolved_path.name}_{uuid.uuid4().hex}.xlsx"
    temp_path = resolved_path.parent / temp_filename

    try:
        wb.save(temp_path)
    finally:
        wb.close()

    # Atomically replace destination
    os.replace(temp_path, resolved_path)

    # 3. Read-back verification
    verify_headers, verify_records = read_spreadsheet_records(resolved_path)
    for k in touched_keys:
        if k not in verify_records:
            raise RuntimeError(
                f"Read-back verification failed: record_key '{k}' was not found after atomic write."
            )

    sha256_hash = compute_file_sha256(resolved_path)

    return UpsertResult(
        inserted=inserted,
        updated=updated,
        unchanged=unchanged,
        keys=touched_keys,
        artifact_path=str(resolved_path),
        sha256=sha256_hash,
    )
