"""Excel (openpyxl) Round-Trip Smoke Test for OFFERHUNT.

Tests creating, writing, saving, and reading back an XLSX file using a temp directory.
"""

import os
import sys
import tempfile
import openpyxl

def main():
    print("=== OFFERHUNT XLSX Smoke Test ===")
    
    with tempfile.TemporaryDirectory() as temp_dir:
        test_file_path = os.path.join(temp_dir, "test_offerhunt.xlsx")
        print(f"Creating test workbook in temporary directory: {temp_dir}")

        # 1. Write workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Offers"

        headers = ["Brand", "Offer Title", "Discount %", "Source URL", "Status"]
        ws.append(headers)

        sample_data = [
            ["Acme Gear", "Summer Flash Sale", 25, "https://example.com/acme", "VERIFIED"],
            ["TechCorp", "Back to School Special", 15, "https://example.com/tech", "PENDING"]
        ]
        for row in sample_data:
            ws.append(row)

        wb.save(test_file_path)
        print(f"Workbook successfully saved to temporary path.")

        # 2. Read back and verify
        loaded_wb = openpyxl.load_workbook(test_file_path)
        loaded_ws = loaded_wb["Offers"]

        # Verify dimensions and content
        loaded_headers = [cell.value for cell in loaded_ws[1]]
        assert loaded_headers == headers, f"Headers mismatch: {loaded_headers} != {headers}"

        rows_read = []
        for row in loaded_ws.iter_rows(min_row=2, max_col=5, values_only=True):
            rows_read.append(list(row))

        assert rows_read == sample_data, f"Data mismatch: {rows_read} != {sample_data}"

        print("\n[PASS] XLSX Round-Trip Successful!")
        print(f"Sheet Name: {loaded_ws.title}")
        print(f"Headers Verified: {len(headers)} columns ({', '.join(headers)})")
        print(f"Rows Verified: {len(rows_read)} sample data rows")
        print("Temporary files cleaned up automatically.")
        return 0

if __name__ == "__main__":
    sys.exit(main())
